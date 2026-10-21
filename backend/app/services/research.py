"""Research Studio: reference discovery, ranking, citation formatting, drafting and export."""
from __future__ import annotations

import io
import logging
import math
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import httpx

from ..config import get_settings
from ..models import Reference, ResearchProject
from . import llm, nlp

log = logging.getLogger("readvibe.research")

SECTION_TITLES = {
    "abstract": "Abstract",
    "introduction": "Introduction",
    "literature_review": "Literature Review",
    "methodology": "Methodology",
    "results": "Results & Discussion",
    "conclusion": "Conclusion",
}


# ─── Discovery providers ──────────────────────────────────────────────────────
def _clean(text: str | None) -> str | None:
    if not text:
        return None
    return " ".join(re.sub(r"<[^>]+>", " ", text).split()) or None


def _semantic_scholar(query: str, limit: int, client: httpx.Client) -> list[dict]:
    settings = get_settings()
    headers = {"x-api-key": settings.semantic_scholar_api_key} if settings.semantic_scholar_api_key else {}
    r = client.get(
        "https://api.semanticscholar.org/graph/v1/paper/search",
        params={
            "query": query,
            "limit": limit,
            "fields": "title,authors,year,venue,abstract,citationCount,externalIds,url,openAccessPdf,publicationTypes",
        },
        headers=headers,
    )
    r.raise_for_status()
    out = []
    for p in r.json().get("data", []) or []:
        if not p.get("title"):
            continue
        ext = p.get("externalIds") or {}
        pdf = (p.get("openAccessPdf") or {}).get("url")
        out.append({
            "title": p["title"],
            "authors": [a.get("name") for a in (p.get("authors") or []) if a.get("name")][:8],
            "year": p.get("year"),
            "venue": p.get("venue") or None,
            "abstract": _clean(p.get("abstract")),
            "url": pdf or p.get("url"),
            "doi": ext.get("DOI"),
            "citations": p.get("citationCount") or 0,
            "provider": "Semantic Scholar",
            "publication_type": ((p.get("publicationTypes") or ["Article"])[0] or "Article").replace("JournalArticle", "Journal article"),
            "open_access": bool(pdf),
        })
    return out


def _crossref(query: str, limit: int, client: httpx.Client) -> list[dict]:
    settings = get_settings()
    r = client.get(
        "https://api.crossref.org/works",
        params={"query": query, "rows": limit, "select": "title,author,issued,container-title,abstract,DOI,URL,is-referenced-by-count,type,license"},
        headers={"User-Agent": f"ReadVibe/3.0 (mailto:{settings.contact_email})"},
    )
    r.raise_for_status()
    out = []
    for item in r.json().get("message", {}).get("items", []):
        title = (item.get("title") or [None])[0]
        if not title:
            continue
        year = None
        parts = (item.get("issued") or {}).get("date-parts") or [[None]]
        if parts and parts[0] and parts[0][0]:
            year = int(parts[0][0])
        authors = [
            f"{a.get('given', '')} {a.get('family', '')}".strip()
            for a in item.get("author", []) if a.get("family")
        ][:8]
        out.append({
            "title": _clean(title),
            "authors": authors,
            "year": year,
            "venue": (item.get("container-title") or [None])[0],
            "abstract": _clean(item.get("abstract")),
            "url": item.get("URL"),
            "doi": item.get("DOI"),
            "citations": item.get("is-referenced-by-count") or 0,
            "provider": "Crossref",
            "publication_type": (item.get("type") or "article").replace("-", " ").capitalize(),
            "open_access": bool(item.get("license")),
        })
    return out


def _openalex_abstract(inv: dict | None) -> str | None:
    if not inv:
        return None
    positions = [(pos, word) for word, poss in inv.items() for pos in poss]
    return " ".join(w for _, w in sorted(positions))


def _openalex(query: str, limit: int, client: httpx.Client) -> list[dict]:
    settings = get_settings()
    r = client.get(
        "https://api.openalex.org/works",
        params={"search": query, "per-page": limit, "mailto": settings.contact_email},
    )
    r.raise_for_status()
    out = []
    for w in r.json().get("results", []):
        if not w.get("title"):
            continue
        oa = w.get("open_access") or {}
        loc = w.get("primary_location") or {}
        out.append({
            "title": _clean(w["title"]),
            "authors": [a["author"]["display_name"] for a in w.get("authorships", []) if a.get("author")][:8],
            "year": w.get("publication_year"),
            "venue": ((loc.get("source") or {}) or {}).get("display_name"),
            "abstract": _clean(_openalex_abstract(w.get("abstract_inverted_index"))),
            "url": oa.get("oa_url") or w.get("doi") or w.get("id"),
            "doi": (w.get("doi") or "").replace("https://doi.org/", "") or None,
            "citations": w.get("cited_by_count") or 0,
            "provider": "OpenAlex",
            "publication_type": (w.get("type") or "article").capitalize(),
            "open_access": bool(oa.get("is_oa")),
        })
    return out


PROVIDERS = {"Semantic Scholar": _semantic_scholar, "Crossref": _crossref, "OpenAlex": _openalex}


def discover(query: str, limit: int = 12, local_sources: list[tuple[str, str, str]] | None = None) -> tuple[list[dict], list[str]]:
    """Query every provider in parallel, de-duplicate and rank. Returns (results, warnings)."""
    settings = get_settings()
    warnings: list[str] = []
    results: list[dict] = []
    if settings.research_online:
        per = max(5, limit)
        with httpx.Client(timeout=settings.research_timeout_seconds, follow_redirects=True) as client:
            with ThreadPoolExecutor(max_workers=len(PROVIDERS)) as pool:
                futures = {name: pool.submit(fn, query, per, client) for name, fn in PROVIDERS.items()}
                for name, fut in futures.items():
                    try:
                        results.extend(fut.result())
                    except Exception as exc:  # noqa: BLE001
                        log.info("%s failed: %s", name, exc)
                        warnings.append(f"{name} is unavailable right now")
    else:
        warnings.append("Online discovery is disabled; showing your own documents only")

    for doc_id, title, author in local_sources or []:
        results.append({
            "title": title, "authors": [author] if author else [], "year": datetime.now().year, "venue": "Your library",
            "abstract": None, "url": None, "doi": None, "citations": 0, "provider": "Your library",
            "publication_type": "Uploaded source", "open_access": True, "local_document_id": doc_id,
        })
    ranked = rank(query, _dedupe(results))
    if not ranked and not warnings:
        warnings.append("No papers matched — try broader keywords")
    return ranked[:limit], warnings


def _norm_title(t: str) -> str:
    return re.sub(r"[^a-z0-9]", "", t.lower())[:80]


def _dedupe(items: list[dict]) -> list[dict]:
    merged: dict[str, dict] = {}
    for it in items:
        key = (it.get("doi") or "").lower() or _norm_title(it["title"])
        alt = _norm_title(it["title"])
        existing = merged.get(key) or next((v for v in merged.values() if _norm_title(v["title"]) == alt), None)
        if existing:
            for f in ("abstract", "url", "doi", "venue", "year"):
                existing[f] = existing.get(f) or it.get(f)
            existing["citations"] = max(existing.get("citations") or 0, it.get("citations") or 0)
            existing["open_access"] = existing.get("open_access") or it.get("open_access")
            if it["provider"] not in existing["provider"]:
                existing["provider"] += f" + {it['provider']}"
        else:
            merged[key] = dict(it)
    return list(merged.values())


def rank(query: str, items: list[dict]) -> list[dict]:
    q = {nlp.stem(w) for w in nlp.content_words(query)}
    year_now = datetime.now().year
    for it in items:
        title_terms = {nlp.stem(w) for w in nlp.content_words(it["title"])}
        abs_terms = {nlp.stem(w) for w in nlp.content_words(it.get("abstract") or "")}
        title_match = len(q & title_terms) / len(q) if q else 0
        abs_match = len(q & abs_terms) / len(q) if q else 0
        cites = math.log10((it.get("citations") or 0) + 1) / 4
        recency = 0.0
        if it.get("year"):
            recency = max(0.0, 1 - (year_now - int(it["year"])) / 25)
        score = 0.5 * title_match + 0.25 * abs_match + 0.15 * min(cites, 1) + 0.1 * recency
        if it.get("provider") == "Your library":
            score = max(score, 0.6)
        it["relevance"] = round(min(0.99, 0.35 + score * 0.7), 3)
        it["summary"] = _relevance_summary(it, q)
    return sorted(items, key=lambda i: i["relevance"], reverse=True)


def _relevance_summary(it: dict, q: set[str]) -> str:
    abstract = it.get("abstract")
    if abstract:
        best = nlp.top_sentences(abstract, 1, query=" ".join(q))
        if best:
            return best[0][:300]
    if it.get("provider") == "Your library":
        return "One of your uploaded documents."
    return "No abstract available — open the source to judge its relevance."


# ─── Citation formatting ──────────────────────────────────────────────────────
def _apa_author(name: str) -> str:
    parts = name.replace(",", "").split()
    if len(parts) == 1:
        return parts[0]
    last = parts[-1]
    initials = " ".join(f"{p[0]}." for p in parts[:-1] if p)
    return f"{last}, {initials}"


def _ieee_author(name: str) -> str:
    parts = name.replace(",", "").split()
    if len(parts) == 1:
        return parts[0]
    return " ".join(f"{p[0]}." for p in parts[:-1]) + " " + parts[-1]


def format_citation(ref: Reference, style: str, number: int) -> str:
    authors = ref.authors or ["Anonymous"]
    year = ref.year or "n.d."
    venue = ref.venue or ""
    doi = f" https://doi.org/{ref.doi}" if ref.doi else (f" {ref.url}" if ref.url else "")
    if style == "IEEE":
        names = ", ".join(_ieee_author(a) for a in authors[:6]) + (" et al." if len(authors) > 6 else "")
        return f"[{number}] {names}, “{ref.title},” {f'{venue}, ' if venue else ''}{year}.{doi}"
    if style == "MLA":
        first = authors[0]
        p = first.split()
        lead = f"{p[-1]}, {' '.join(p[:-1])}" if len(p) > 1 else first
        more = ", et al" if len(authors) > 2 else (f", and {authors[1]}" if len(authors) == 2 else "")
        return f"{lead}{more}. “{ref.title}.” {f'{venue}, ' if venue else ''}{year}.{doi}"
    # APA 7
    if len(authors) == 1:
        names = _apa_author(authors[0])
    elif len(authors) <= 20:
        names = ", ".join(_apa_author(a) for a in authors[:-1]) + ", & " + _apa_author(authors[-1])
    else:
        names = ", ".join(_apa_author(a) for a in authors[:19]) + ", … " + _apa_author(authors[-1])
    return f"{names} ({year}). {ref.title}. {f'{venue}. ' if venue else ''}{doi.strip()}".strip()


def in_text(ref: Reference, style: str, number: int) -> str:
    if style == "IEEE":
        return f"[{number}]"
    last = (ref.authors[0].split()[-1] if ref.authors else "Anonymous")
    suffix = " et al." if len(ref.authors or []) > 2 else (f" & {ref.authors[1].split()[-1]}" if len(ref.authors or []) == 2 else "")
    return f"({last}{suffix}, {ref.year or 'n.d.'})" if style == "APA" else f"({last}{suffix})"


def citation_warnings(project: ResearchProject) -> list[dict]:
    warnings = []
    included = [r for r in project.references if r.included]
    if not included:
        warnings.append({"level": "error", "message": "No references are included — the paper has no supporting sources."})
    for r in included:
        missing = [f for f, v in (("authors", r.authors), ("year", r.year), ("venue", r.venue), ("DOI or URL", r.doi or r.url)) if not v]
        if missing:
            warnings.append({"level": "warning", "message": f"“{r.title[:70]}” is missing {', '.join(missing)}.", "reference_id": r.id})
    numbers = {str(i + 1) for i in range(len(included))}
    full = " ".join((s or {}).get("content", "") for s in (project.sections or {}).values())
    cited = set(re.findall(r"\[(\d+)\]", full))
    for n in sorted(cited - numbers, key=int):
        warnings.append({"level": "error", "message": f"Citation [{n}] does not match any included reference."})
    unused = numbers - cited
    if full.strip() and unused:
        warnings.append({"level": "info", "message": f"{len(unused)} included reference(s) are never cited in the draft."})
    return warnings


# ─── Drafting ─────────────────────────────────────────────────────────────────
def _ref_context(project: ResearchProject) -> tuple[str, list[Reference]]:
    included = [r for r in project.references if r.included]
    lines = []
    for i, r in enumerate(included, 1):
        lines.append(f"[{i}] {r.title} ({', '.join(r.authors[:3])}{' et al.' if len(r.authors) > 3 else ''}, {r.year or 'n.d.'}). "
                     f"Abstract: {r.abstract or r.summary or 'not available'}")
    return "\n".join(lines), included


def draft_section(project: ResearchProject, key: str, action: str, source_text: str = "") -> str:
    current = ((project.sections or {}).get(key) or {}).get("content", "")
    refs_text, included = _ref_context(project)
    title = SECTION_TITLES[key]
    if llm.available():
        instructions = {
            "generate": f"Write the {title} section.",
            "regenerate": f"Write a fresh alternative version of the {title} section.",
            "expand": "Expand the current draft with more depth and supporting evidence from the references.",
            "shorten": "Shorten the current draft to roughly half its length while keeping the key claims and citations.",
            "academic": "Rewrite the current draft in a more formal academic tone without changing its claims.",
            "cite": "Add citations from the reference list to every claim in the current draft that needs support.",
            "simplify": "Rewrite the current draft in clearer, simpler language while keeping citations.",
        }[action]
        prompt = (
            f"Research topic: {project.topic}\nPaper title: {project.title}\nTone: {project.tone}. Depth: {project.depth}.\n"
            f"{instructions}\nReturn only the section body text (no heading), in paragraphs.\n\n"
            f"<references>\n{refs_text or 'No references selected.'}\n</references>\n"
            + (f"<uploaded_sources>\n{source_text[:12000]}\n</uploaded_sources>\n" if source_text else "")
            + (f"<current_draft>\n{current}\n</current_draft>" if current and action not in ("generate",) else "")
        )
        try:
            return llm.complete_text(llm.RESEARCH_SYSTEM, [{"role": "user", "content": prompt}], max_tokens=8000)
        except llm.LLMUnavailable as exc:
            log.info("draft fallback: %s", exc)
    return _local_draft(project, key, action, current, included, source_text)


def _local_draft(project: ResearchProject, key: str, action: str, current: str, refs: list[Reference], source_text: str) -> str:
    if action in ("expand", "shorten", "academic", "cite", "simplify") and current.strip():
        return _transform(current, action, refs, project)
    topic = project.topic.rstrip(".")
    evidence = []
    for i, r in enumerate(refs, 1):
        basis = r.abstract or r.summary
        if basis:
            best = nlp.top_sentences(basis, 1, query=topic)
            if best:
                evidence.append((i, r, _paraphrase(best[0])))
    src_points = nlp.top_sentences(source_text, 3, query=topic) if source_text else []
    years = sorted({r.year for r in refs if r.year})
    span = f"between {years[0]} and {years[-1]}" if len(years) > 1 else (f"in {years[0]}" if years else "")
    cite_all = ", ".join(f"[{i}]" for i in range(1, min(len(refs), 4) + 1))
    concepts = nlp.key_phrases(" ".join(r.title for r in refs) + " " + topic, 5)
    concept_text = ", ".join(c.lower() for c in concepts[:4]) or topic.lower()

    if key == "abstract":
        parts = [
            f"This paper examines {topic.lower()}.",
            f"Drawing on {len(refs)} sources published {span}, it synthesises current understanding of {concept_text}." if refs else "",
            f"Prior work indicates that {evidence[0][2][0].lower() + evidence[0][2][1:]} [{evidence[0][0]}]." if evidence else "",
            "The paper outlines the methodology used to review the literature, discusses the principal findings, and identifies open questions for future research.",
        ]
        return " ".join(p for p in parts if p)
    if key == "introduction":
        return "\n\n".join(p for p in [
            f"{topic} has attracted growing attention{(' ' + cite_all) if refs else ''}. Understanding it matters because it shapes both research practice and real-world outcomes.",
            " ".join(f"{e[2]} [{e[0]}]." for e in evidence[:2]),
            f"This paper addresses the following question: what does current research reveal about {topic.lower()}, and where are the gaps? "
            "The remainder of the paper reviews the literature, describes the approach, discusses the results and concludes with implications.",
        ] if p.strip())
    if key == "literature_review":
        if not evidence and not src_points:
            return "Add or include references with abstracts to generate a literature review grounded in your sources."
        paras = [f"Research on {topic.lower()} spans {len(refs)} key sources {span}. The following themes emerge."]
        for c in concepts[:3]:
            related = [e for e in evidence if c.lower().split()[0] in e[2].lower() or c.lower().split()[0] in e[1].title.lower()]
            if related:
                paras.append(f"Regarding {c.lower()}, " + " ".join(f"{e[2]} [{e[0]}]." for e in related[:3]))
        remaining = [e for e in evidence if not any(f"[{e[0]}]" in p for p in paras)]
        if remaining:
            paras.append("Further studies add to this picture. " + " ".join(f"{e[2]} [{e[0]}]." for e in remaining[:4]))
        if src_points:
            paras.append("The uploaded source material similarly notes that " + " ".join(src_points))
        paras.append("Taken together, these works suggest areas of agreement as well as open questions that motivate the present study.")
        return "\n\n".join(paras)
    if key == "methodology":
        return "\n\n".join([
            f"A structured literature review was conducted on {topic.lower()}. Candidate sources were retrieved from "
            f"{', '.join(sorted({r.provider.split(' + ')[0] for r in refs})) or 'academic databases'} using topic keywords "
            f"({concept_text}).",
            f"Sources were screened for relevance using title and abstract matching, citation impact and recency; "
            f"{len(refs)} were included for synthesis. Each was coded for its central claims, methods and findings, "
            "and themes were identified by comparing these codes across studies.",
            "Limitations of this approach include reliance on indexed databases and abstracts, which may under-represent grey literature.",
        ])
    if key == "results":
        if not evidence:
            return "Include references with abstracts to synthesise results."
        return "\n\n".join([
            "The review yields several findings.",
            f"First, {evidence[0][2][0].lower() + evidence[0][2][1:]} [{evidence[0][0]}].",
            " ".join(f"In addition, {e[2][0].lower() + e[2][1:]} [{e[0]}]." for e in evidence[1:4]),
            "These results should be interpreted with caution: the included studies differ in context and method, and further empirical work is needed to confirm the patterns observed.",
        ])
    # conclusion
    return "\n\n".join([
        f"This paper reviewed research on {topic.lower()}.",
        f"Across {len(refs)} sources, the evidence highlights {concept_text}{(' ' + cite_all) if refs else ''}.",
        "Future work should test these findings with larger and more diverse samples, and examine how they translate into practice.",
    ])


def _paraphrase(sentence: str) -> str:
    s = nlp.simplify_sentence(sentence)
    s = re.sub(r"^(In this (paper|study|article|work),?\s*)?(we|the authors)\s+", "the authors ", s, flags=re.I)
    s = re.sub(r"\b(our|we)\b", "their", s, flags=re.I)
    s = s.rstrip(".")
    return s[0].upper() + s[1:] if s else s


def _transform(text: str, action: str, refs: list[Reference], project: ResearchProject) -> str:
    paras = nlp.paragraphs(text)
    if action == "shorten":
        keep = max(2, len(nlp.sentences(text)) // 2)
        return " ".join(nlp.top_sentences(text, keep))
    if action == "simplify":
        return "\n\n".join(" ".join(nlp.simplify_sentence(s) for s in nlp.sentences(p)) for p in paras)
    if action == "academic":
        repl = {r"\bshows\b": "demonstrates", r"\bshow\b": "demonstrate", r"\bbig\b": "substantial", r"\blots of\b": "considerable",
                r"\bget\b": "obtain", r"\bhelp\b": "facilitate", r"\bthink\b": "posit", r"\bdon't\b": "do not", r"\bcan't\b": "cannot",
                r"\bit's\b": "it is", r"\bkind of\b": "somewhat", r"\bso\b": "consequently", r"\bbut\b": "however,", r"\balso\b": "furthermore,"}
        out = text
        for p, r in repl.items():
            out = re.sub(p, r, out)
        return out
    if action == "cite":
        if not refs:
            return text
        out_paras = []
        for p in paras:
            sents = nlp.sentences(p)
            new = []
            for s in sents:
                if re.search(r"\[\d+\]", s) or len(s.split()) < 8:
                    new.append(s)
                    continue
                best_i, best = 0, 0
                for i, r in enumerate(refs, 1):
                    overlap = len(set(nlp.content_words(s)) & set(nlp.content_words(r.title + " " + (r.abstract or ""))))
                    if overlap > best:
                        best_i, best = i, overlap
                new.append(s.rstrip(".") + f" [{best_i}]." if best >= 2 else s)
            out_paras.append(" ".join(new))
        return "\n\n".join(out_paras)
    # expand
    _, included = _ref_context(project)
    extra = []
    for i, r in enumerate(included, 1):
        if r.abstract and f"[{i}]" not in text:
            best = nlp.top_sentences(r.abstract, 1, query=project.topic)
            if best:
                extra.append(f"{_paraphrase(best[0])} [{i}].")
        if len(extra) >= 3:
            break
    addition = " ".join(extra) or "Further evidence from additional sources would strengthen this section."
    return text.rstrip() + "\n\n" + "Building on this, " + addition[0].lower() + addition[1:]


# ─── Export ───────────────────────────────────────────────────────────────────
def _ordered_sections(project: ResearchProject) -> list[tuple[str, str]]:
    sections = project.sections or {}
    return [(SECTION_TITLES[k], (sections.get(k) or {}).get("content", "").strip()) for k in SECTION_TITLES]


def _bibliography(project: ResearchProject) -> list[str]:
    included = [r for r in project.references if r.included]
    return [format_citation(r, project.citation_style, i) for i, r in enumerate(included, 1)]


def export_markdown(project: ResearchProject, author: str) -> bytes:
    lines = [f"# {project.title}", "", f"*{author} · {datetime.now():%B %d, %Y}*", ""]
    for title, body in _ordered_sections(project):
        if body:
            lines += [f"## {title}", "", body, ""]
    bib = _bibliography(project)
    if bib:
        lines += ["## References", ""] + [f"- {b}" if project.citation_style != "IEEE" else b for b in bib]
    return "\n".join(lines).encode("utf-8")


def export_docx(project: ResearchProject, author: str) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    title = d.add_heading(project.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = d.add_paragraph(f"{author} · {datetime.now():%B %d, %Y}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in _ordered_sections(project):
        if not body:
            continue
        d.add_heading(heading, level=1)
        for para in nlp.paragraphs(body):
            p = d.add_paragraph(para)
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing = 1.5
    bib = _bibliography(project)
    if bib:
        d.add_heading("References", level=1)
        for b in bib:
            p = d.add_paragraph(b)
            p.paragraph_format.left_indent = Pt(24)
            p.paragraph_format.first_line_indent = Pt(-24)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _latin1(text: str) -> str:
    table = {"“": '"', "”": '"', "‘": "'", "’": "'", "—": "-", "–": "-", "…": "...", "•": "-", "·": "-", "→": "->"}
    for k, v in table.items():
        text = text.replace(k, v)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def export_pdf(project: ResearchProject, author: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(22, 22, 22)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()
    pdf.set_font("Times", "B", 20)
    pdf.multi_cell(0, 10, _latin1(project.title), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Times", "I", 11)
    pdf.cell(0, 8, _latin1(f"{author} - {datetime.now():%B %d, %Y}"), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)
    for heading, body in _ordered_sections(project):
        if not body:
            continue
        pdf.set_font("Times", "B", 14)
        pdf.multi_cell(0, 8, _latin1(heading), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        pdf.set_font("Times", "", 12)
        for para in nlp.paragraphs(body):
            pdf.multi_cell(0, 6.5, _latin1(para), align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2.5)
        pdf.ln(3)
    bib = _bibliography(project)
    if bib:
        pdf.set_font("Times", "B", 14)
        pdf.multi_cell(0, 8, "References", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Times", "", 11)
        for b in bib:
            pdf.multi_cell(0, 6, _latin1(b), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
    pdf.set_font("Times", "I", 8)
    return bytes(pdf.output())
