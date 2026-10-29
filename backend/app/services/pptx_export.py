"""Deck → .pptx (16:9) with native charts, shape-built diagrams, generated images, speaker notes, transitions and alt text."""
from __future__ import annotations

import io
import math
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

from . import art

W, H = 13.333, 7.5
MX = 0.7  # side margin
FONT = "Calibri"


def _c(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


class _Ctx:
    def __init__(self, theme: str, deck_title: str):
        self.th = art.theme_of(theme)
        self.theme = theme
        self.deck_title = deck_title

    def col(self, key: str) -> RGBColor:
        return _c(self.th[key])


# ─── Primitives ───────────────────────────────────────────────────────────────
def _text(slide, x, y, w, h, text, size, color, *, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, FONT
    f.color.rgb = color
    return tb


def _runs(text: str, highlights: list[str]) -> list[tuple[str, bool]]:
    """Split `text` into (piece, highlighted) so key phrases can be emphasised; case-insensitive, first occurrence of each."""
    spans = []
    for h in highlights or []:
        m = re.search(rf"(?<!\w){re.escape(h)}(?!\w)", text, re.I)  # whole words only
        if m and not any(m.start() < b and m.end() > a for a, b in spans):
            spans.append((m.start(), m.end()))
    spans.sort()
    out, at = [], 0
    for a, b in spans:
        if a > at:
            out.append((text[at:a], False))
        out.append((text[a:b], True))
        at = b
    if at < len(text):
        out.append((text[at:], False))
    return out or [(text, False)]


def _bullets(slide, x, y, w, h, items, size, color, bullet_color, numbered=False, highlights=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.55)
        mark = p.add_run()
        mark.text = f"{i + 1}  " if numbered else "●  "
        mark.font.size, mark.font.name, mark.font.bold = Pt(size * (1.0 if numbered else 0.6)), FONT, True
        mark.font.color.rgb = bullet_color
        for piece, hot in _runs(it, highlights or []):
            r = p.add_run()
            r.text = piece
            r.font.size, r.font.name = Pt(size), FONT
            r.font.bold = hot
            r.font.color.rgb = bullet_color if hot else color
    return tb


def _shape(slide, kind, x, y, w, h, fill, line=None, alpha: int | None = None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        if alpha is not None:
            clr = s._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            clr.append(parse_xml(f'<a:alpha {nsdecls("a")} val="{alpha * 1000}"/>'))
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    return s


def _label(shape, text, size, color, *, bold=True, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = color


def _line(slide, x1, y1, x2, y2, color, width=2.0, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        c.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
    return c


def _picture(slide, seed: str, style: str, ctx: _Ctx, x, y, w, h, alt: str):
    """Generated art, cropped (not stretched) to the w×h frame."""
    png = art.render(seed, ctx.theme, style or "", 960, 540)
    pic = slide.shapes.add_picture(io.BytesIO(png), Inches(x), Inches(y), Inches(w), Inches(h))
    target, native = w / h, 960 / 540
    if target < native:
        crop = (1 - target / native) / 2
        pic.crop_left = pic.crop_right = crop
    elif target > native:
        crop = (1 - native / target) / 2
        pic.crop_top = pic.crop_bottom = crop
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def _transition(slide, kind: str):
    inner = {"fade": "<p:fade/>", "push": '<p:push dir="l"/>', "zoom": '<p:zoom dir="in"/>'}.get(kind, "<p:fade/>")
    el = parse_xml(f'<p:transition {nsdecls("p")} spd="med">{inner}</p:transition>')
    anchor = slide._element.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = slide._element.find(qn("p:cSld"))
    anchor.addnext(el)


def _frame(slide, ctx: _Ctx, s: dict, n: int, total: int, *, title=True):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = ctx.col("bg0")
    if title:
        size = 34 if len(s["title"]) <= 46 else 28
        _text(slide, MX, 0.42, W - 2 * MX, 1.0, s["title"], size, ctx.col("text"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _shape(slide, MSO_SHAPE.RECTANGLE, MX + 0.05, 1.42, 0.9, 0.06, ctx.col("a1"))
    _text(slide, MX, H - 0.5, 8, 0.3, ctx.deck_title, 10, ctx.col("muted"))
    _text(slide, W - MX - 1.5, H - 0.5, 1.5, 0.3, f"{n} / {total}", 10, ctx.col("muted"), align=PP_ALIGN.RIGHT)


# ─── Layouts ──────────────────────────────────────────────────────────────────
def _title_slide(slide, ctx, s):
    img = s.get("image") or {"seed": s["title"], "style": ""}
    _picture(slide, img["seed"], img["style"], ctx, 0, 0, W, H, s["alt"])
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 1.9, 7.6, 3.5, ctx.col("bg0"), alpha=78)
    _text(slide, 1.0, 2.15, 7.0, 2.0, s["title"], 40 if len(s["title"]) < 40 else 32, ctx.col("text"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _shape(slide, MSO_SHAPE.RECTANGLE, 1.05, 4.2, 1.1, 0.07, ctx.col("a1"))
    if s.get("subtitle"):
        _text(slide, 1.0, 4.4, 7.0, 0.8, s["subtitle"], 18, ctx.col("muted"))


_CALLOUT_LABEL = {"note": "NOTE", "wow": "WOW!", "tip": "TIP", "remember": "REMEMBER"}
_CALLOUT_COLOR = {"note": "a3", "wow": "a1", "tip": "a2", "remember": "a1"}


def _key_point(slide, ctx, s, y=1.62):
    """The one thing to remember, as an accent pill under the title."""
    if not s.get("main_point"):
        return False
    box = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, MX, y, W - 2 * MX, 0.62, ctx.col("panel"), line=ctx.col("a1"))
    _text(slide, MX + 0.2, y + 0.04, 1.6, 0.54, "KEY POINT", 12, ctx.col("a1"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, MX + 1.8, y + 0.04, W - 2 * MX - 2.0, 0.54, s["main_point"], 18, ctx.col("text"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    return box is not None


def _callouts(slide, ctx, s, y=6.05):
    """Margin notes along the bottom: a notation, an exclamation, a tip or a thing to remember."""
    items = s.get("callouts") or []
    if not items:
        return
    gap = 0.3
    w = (W - 2 * MX - gap * (len(items) - 1)) / len(items)
    for i, c in enumerate(items):
        x = MX + i * (w + gap)
        accent = ctx.col(_CALLOUT_COLOR.get(c["kind"], "a1"))
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.8, ctx.col("panel"), line=accent)
        _text(slide, x + 0.15, y + 0.05, 1.3, 0.7, _CALLOUT_LABEL.get(c["kind"], "NOTE"), 12, accent, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, x + 1.4, y + 0.05, w - 1.55, 0.7, c["text"], 14, ctx.col("text"), anchor=MSO_ANCHOR.MIDDLE)


def _text_slide(slide, ctx, s, numbered=False):
    items = s["bullets"]
    img = s.get("image")
    key = _key_point(slide, ctx, s)
    has_notes = bool(s.get("callouts"))
    top = 2.5 if key else 1.85
    room = (5.9 if has_notes else 6.7) - top
    size = 24 if len(items) <= 4 and not (key and has_notes) else 20
    width = 7.2 if img else W - 2 * MX
    _bullets(slide, MX, top, width, room, items, size, ctx.col("text"), ctx.col("a1"), numbered=numbered, highlights=s.get("highlights"))
    if img:
        _picture(slide, img["seed"], img["style"], ctx, 8.35, top + 0.1, 4.3, min(3.9, room - 0.1), s["alt"])
    _callouts(slide, ctx, s)


def _example_slide(slide, ctx, s):
    """An example: its generated illustration on the left, the example and what it shows on the right."""
    d = s.get("data") or {}
    img = s.get("image") or {"seed": d.get("text") or s["title"], "style": "scene"}
    _picture(slide, img["seed"], img["style"], ctx, MX, 1.85, 5.4, 4.0, s["alt"])
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.5, 1.85, 1.5, 0.42, ctx.col("a1"))
    _text(slide, 6.5, 1.85, 1.5, 0.42, "EXAMPLE", 12, ctx.col("bg0"), bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, 6.5, 2.4, W - 6.5 - MX, 2.4, d.get("text", ""), 20, ctx.col("text"))
    if d.get("takeaway"):
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 6.5, 4.95, W - 6.5 - MX, 0.9, ctx.col("panel"), line=ctx.col("a2"))
        _text(slide, 6.7, 5.0, 1.4, 0.8, "SO WHAT", 12, ctx.col("a2"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, 8.0, 5.0, W - 8.0 - MX - 0.2, 0.8, d["takeaway"], 16, ctx.col("text"), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _callouts(slide, ctx, s)


def _quote_slide(slide, ctx, s):
    img = s.get("image") or {"seed": s["title"], "style": ""}
    _picture(slide, img["seed"], img["style"], ctx, 0, 0, W, H, s["alt"])
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.2, 1.5, W - 2.4, 4.5, ctx.col("bg0"), alpha=82)
    _text(slide, 1.6, 1.7, 1.2, 1.2, "“", 80, ctx.col("a1"), bold=True)
    d = s["data"]
    _text(slide, 1.9, 2.5, W - 3.8, 2.6, d["text"], 30 if len(d["text"]) < 120 else 25, ctx.col("text"), italic=True, anchor=MSO_ANCHOR.MIDDLE)
    if d.get("source"):
        _text(slide, 1.9, 5.15, W - 3.8, 0.5, "— " + d["source"], 16, ctx.col("muted"))


def _chart_slide(slide, ctx, s):
    d = s["data"]
    cd = CategoryChartData()
    cd.categories = d["labels"]
    cd.add_series(d.get("unit") or "Value", d["values"])
    kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "donut": XL_CHART_TYPE.DOUGHNUT}[d["kind"]]
    gf = slide.shapes.add_chart(kind, Inches(MX), Inches(1.7), Inches(8.6), Inches(4.9), cd)
    gf._element.nvGraphicFramePr.cNvPr.set("descr", s["alt"])
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(13), FONT
    ch.font.color.rgb = ctx.col("text")
    ch.has_title = False
    palette = [ctx.col("a1"), ctx.col("a2"), ctx.col("a3")]
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(12)
    plot.data_labels.font.color.rgb = ctx.col("text")
    if d["kind"] == "donut":
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = palette[i % 3]
    else:
        ch.has_legend = False
        ser = plot.series[0]
        if d["kind"] == "bar":
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = palette[0]
            plot.gap_width = 60
        else:
            ser.format.line.color.rgb = palette[0]
            ser.format.line.width = Pt(3)
        ch.value_axis.major_gridlines.format.line.color.rgb = ctx.col("panel")
        ch.value_axis.format.line.fill.background()
        ch.category_axis.format.line.color.rgb = ctx.col("muted")
    cap = d.get("caption") or ""
    _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 9.6, 1.9, 3.0, 4.4, ctx.col("panel"))
    _text(slide, 9.8, 2.05, 2.6, 0.5, "WHAT IT SHOWS", 11, ctx.col("a1"), bold=True)
    _text(slide, 9.8, 2.5, 2.6, 3.6, cap or s["title"], 16, ctx.col("text"))


def _process_slide(slide, ctx, s):
    d = s["data"]
    nodes = d["nodes"]
    per_row = 4 if len(nodes) > 4 else len(nodes)
    rows = math.ceil(len(nodes) / per_row)
    gap, bw, bh = 0.55, (W - 2 * MX - 0.55 * (per_row - 1)) / per_row, 1.35
    top = 2.2 if rows == 1 else 1.95
    pos = []
    for i, n in enumerate(nodes):
        r, c = divmod(i, per_row)
        x = MX + c * (bw + gap)
        y = top + r * (bh + 1.1)
        box = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh, ctx.col("panel"), ctx.col("a1"))
        _label(box, n["label"], 17, ctx.col("text"))
        badge = _shape(slide, MSO_SHAPE.OVAL, x - 0.15, y - 0.17, 0.45, 0.45, ctx.col("a1"))
        _label(badge, str(i + 1), 12, ctx.col("bg0"))
        if n.get("detail"):
            _text(slide, x, y + bh + 0.05, bw, 0.7, n["detail"], 12, ctx.col("muted"), align=PP_ALIGN.CENTER)
        pos.append((x, y))
    for i in range(len(nodes) - 1):
        (x1, y1), (x2, y2) = pos[i], pos[i + 1]
        if y1 == y2:
            _line(slide, x1 + bw, y1 + bh / 2, x2, y2 + bh / 2, ctx.col("a2"), 2.5, arrow=True)
        else:  # wrap to the next row
            _line(slide, x1 + bw / 2, y1 + bh, x2 + bw / 2, y2, ctx.col("a2"), 2.5, arrow=True)
    if d.get("loop"):
        _text(slide, MX, H - 1.05, W - 2 * MX, 0.4, "↻  the process repeats from step 1", 14, ctx.col("a3"), align=PP_ALIGN.CENTER)


def _hierarchy_slide(slide, ctx, s):
    root = s["data"]["root"]
    kids = root["children"]
    rw, rh = 3.4, 0.85
    rx = (W - rw) / 2
    ry = 1.8
    top = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, rx, ry, rw, rh, ctx.col("a1"))
    _label(top, root["label"], 18, ctx.col("bg0"))
    k = len(kids)
    cw = (W - 2 * MX) / k
    bw = min(2.3, cw - 0.25)
    for i, kid in enumerate(kids):
        cx = MX + cw * (i + 0.5)
        y = 3.2
        _line(slide, W / 2, ry + rh, cx, y, ctx.col("a2"), 2)
        box = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - bw / 2, y, bw, 0.75, ctx.col("panel"), ctx.col("a2"))
        _label(box, kid["label"], 15, ctx.col("text"))
        for j, g in enumerate(kid["children"][:4]):
            gy = y + 1.1 + j * 0.62
            _line(slide, cx, y + 0.75 if j == 0 else gy - 0.1, cx, gy, ctx.col("muted"), 1.25)
            gb = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - bw / 2 + 0.1, gy, bw - 0.2, 0.48, ctx.col("bg1"), ctx.col("muted"))
            _label(gb, g["label"], 12, ctx.col("text"), bold=False)


def _mindmap_slide(slide, ctx, s):
    d = s["data"]
    cx, cy = W / 2, 4.2
    center = _shape(slide, MSO_SHAPE.OVAL, cx - 1.35, cy - 0.65, 2.7, 1.3, ctx.col("a1"))
    _label(center, d.get("root") or s["title"], 17, ctx.col("bg0"))
    n = len(d["branches"])
    rx, ry = 4.6, 2.25
    palette = [ctx.col("a2"), ctx.col("a3"), ctx.col("a1")]
    for i, b in enumerate(d["branches"]):
        ang = -math.pi / 2 + (i + 0.5) * 2 * math.pi / n
        bx, by = cx + rx * math.cos(ang), cy + ry * math.sin(ang)
        col = palette[i % 3]
        _line(slide, cx + 1.35 * math.cos(ang), cy + 0.65 * math.sin(ang), bx, by, col, 2.5)
        node = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx - 1.05, by - 0.3, 2.1, 0.6, ctx.col("panel"), col)
        _label(node, b["label"], 14, ctx.col("text"))
        if b["children"]:
            side = 1 if math.cos(ang) >= 0 else -1
            tx = bx + (1.1 if side > 0 else -2.95)
            _text(slide, tx, by - 0.38, 1.9, 0.8, "\n".join(f"· {c}" for c in b["children"][:3]), 11, ctx.col("muted"),
                  align=PP_ALIGN.LEFT if side > 0 else PP_ALIGN.RIGHT)


def _columns_slide(slide, ctx, s, versus=False):
    cols = s["data"]["columns"]
    k = len(cols)
    gap = 0.35
    cw = (W - 2 * MX - gap * (k - 1)) / k
    palette = [ctx.col("a1"), ctx.col("a2"), ctx.col("a3"), ctx.col("a1")]
    for i, c in enumerate(cols):
        x = MX + i * (cw + gap)
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.9, cw, 4.6, ctx.col("panel"))
        head = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.9, cw, 0.8, palette[i])
        _label(head, c["title"], 18, ctx.col("bg0"))
        _bullets(slide, x + 0.2, 2.95, cw - 0.4, 3.4, c["points"], 15, ctx.col("text"), palette[i])
    if versus and k == 2:
        v = _shape(slide, MSO_SHAPE.OVAL, W / 2 - 0.35, 3.9, 0.7, 0.7, ctx.col("bg0"), ctx.col("a3"))
        _label(v, "VS", 13, ctx.col("a3"))


def _timeline_slide(slide, ctx, s):
    ev = s["data"]["events"]
    n = len(ev)
    y = 4.0
    _line(slide, MX, y, W - MX, y, ctx.col("muted"), 3)
    cw = (W - 2 * MX) / n
    for i, e in enumerate(ev):
        cx = MX + cw * (i + 0.5)
        dot = _shape(slide, MSO_SHAPE.OVAL, cx - 0.17, y - 0.17, 0.34, 0.34, ctx.col("a1"), ctx.col("bg0"))
        up = i % 2 == 0
        _text(slide, cx - cw / 2, y - 1.25 if up else y + 0.3, cw, 0.5, e["label"], 22, ctx.col("a1"), bold=True, align=PP_ALIGN.CENTER)
        _text(slide, cx - cw / 2 + 0.05, y - 0.75 if up else y + 0.8, cw - 0.1, 1.2, e["detail"], 12, ctx.col("text"), align=PP_ALIGN.CENTER)
        del dot


def _stats_slide(slide, ctx, s):
    items = s["data"]["items"]
    k = len(items)
    gap = 0.4
    cw = (W - 2 * MX - gap * (k - 1)) / k
    for i, it in enumerate(items):
        x = MX + i * (cw + gap)
        _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.1, cw, 3.8, ctx.col("panel"))
        _text(slide, x, 2.5, cw, 1.4, it["value"], 48 if len(it["value"]) < 8 else 36, ctx.col("a1"), bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, x + 0.2, 4.1, cw - 0.4, 1.6, it["label"], 15, ctx.col("text"), align=PP_ALIGN.CENTER)


def _add_slide(prs: Presentation, ctx: _Ctx, s: dict, n: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    layout = s["layout"]
    full_bleed = layout in ("title", "quote")
    _frame(slide, ctx, s, n, total, title=not full_bleed)
    if layout == "title":
        _title_slide(slide, ctx, s)
    elif layout == "quote":
        _quote_slide(slide, ctx, s)
    elif layout == "chart":
        _chart_slide(slide, ctx, s)
    elif layout == "process":
        _process_slide(slide, ctx, s)
    elif layout == "hierarchy":
        _hierarchy_slide(slide, ctx, s)
    elif layout == "mindmap":
        _mindmap_slide(slide, ctx, s)
    elif layout in ("types", "compare"):
        _columns_slide(slide, ctx, s, versus=layout == "compare")
    elif layout == "timeline":
        _timeline_slide(slide, ctx, s)
    elif layout == "stats":
        _stats_slide(slide, ctx, s)
    elif layout == "example":
        _example_slide(slide, ctx, s)
    else:  # concept · bullets · agenda · summary
        _text_slide(slide, ctx, s, numbered=layout == "agenda")
    if full_bleed:  # the picture sits above the frame's footer; bring the footer back on top
        _text(slide, MX, H - 0.5, 8, 0.3, ctx.deck_title, 10, ctx.col("muted"))
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]
    _transition(slide, s.get("transition", "fade"))


def build_pptx(deck_title: str, theme: str, slides: list[dict]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(W * 914400)), Emu(int(H * 914400))
    ctx = _Ctx(theme, deck_title)
    prs.core_properties.title = deck_title
    prs.core_properties.author = "ReadVibe"
    for i, s in enumerate(slides):
        _add_slide(prs, ctx, s, i + 1, len(slides))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
