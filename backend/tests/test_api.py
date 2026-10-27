import io

import docx

from app.seed import DESIGN_THINKING, FOCUSED_READING


def _upload_text(client, auth, name="notes.txt", body=FOCUSED_READING):
    r = client.post("/api/documents/upload", headers=auth, files={"file": (name, body.encode(), "text/plain")})
    assert r.status_code == 202, r.text
    return r.json()


# ─── Auth ─────────────────────────────────────────────────────────────────────
def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json()["database"] is True


def test_auth_flow(client):
    r = client.post("/api/auth/register", json={"name": "A", "email": "a@example.com", "password": "password123"})
    assert r.status_code == 201
    assert client.post("/api/auth/register", json={"name": "A", "email": "a@example.com", "password": "password123"}).status_code == 409
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "wrong-pass"}).status_code == 401
    token = client.post("/api/auth/login", json={"email": "A@example.com", "password": "password123"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/auth/me", headers=h).json()["email"] == "a@example.com"
    assert client.put("/api/auth/me/preferences", headers=h, json={"preferences": {"theme": "light"}}).json()["preferences"]["theme"] == "light"
    assert client.post("/api/auth/me/password", headers=h, json={"current_password": "password123", "new_password": "newpassword1"}).status_code == 204
    assert client.post("/api/auth/login", json={"email": "a@example.com", "password": "newpassword1"}).status_code == 200
    assert client.get("/api/auth/me/export", headers=h).status_code == 200
    assert client.delete("/api/auth/me", headers=h).status_code == 204
    assert client.get("/api/auth/me", headers=h).status_code == 401


def test_requires_auth(client):
    assert client.get("/api/documents").status_code == 401


def test_validation_errors_are_readable(client):
    r = client.post("/api/auth/register", json={"name": "x", "email": "bad", "password": "1"})
    assert r.status_code == 422
    assert "detail" in r.json()


# ─── Documents ────────────────────────────────────────────────────────────────
def test_upload_and_read_text(client, auth):
    data = _upload_text(client, auth)
    doc_id = data["document"]["id"]
    job = client.get(f"/api/jobs/{data['job_id']}", headers=auth).json()
    assert job["status"] == "completed", job
    doc = client.get(f"/api/documents/{doc_id}", headers=auth).json()
    assert doc["status"] == "ready"
    assert doc["chapter_count"] == 4
    assert doc["chapters"][1]["title"] == "Building a Mental Model"
    ch = client.get(f"/api/documents/{doc_id}/chapters/1", headers=auth).json()
    assert "mental model" in ch["content"]

    prog = client.put(f"/api/documents/{doc_id}/progress", headers=auth, json={"chapter_index": 2, "scroll_ratio": 0.5, "mode": "Zen"}).json()
    assert 0 < prog["progress"] < 100

    found = client.get(f"/api/documents/{doc_id}/search", headers=auth, params={"q": "retrieval"}).json()
    assert found["total"] >= 2

    upd = client.patch(f"/api/documents/{doc_id}", headers=auth, json={"favorite": True, "tags": ["Study"], "title": "Renamed"}).json()
    assert upd["favorite"] and upd["title"] == "Renamed"
    assert any(d["id"] == doc_id for d in client.get("/api/documents", headers=auth, params={"favorites": True}).json())
    assert client.get(f"/api/documents/{doc_id}/package", headers=auth).json()["chapters"]


def test_upload_docx(client, auth):
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("Design thinking is a process for creative problem solving. " * 20)
    d.add_heading("Methods", level=1)
    d.add_paragraph("We interviewed users and built prototypes to test our ideas. " * 20)
    buf = io.BytesIO()
    d.save(buf)
    r = client.post("/api/documents/upload", headers=auth, files={"file": ("study.docx", buf.getvalue(), "application/octet-stream")})
    doc = client.get(f"/api/documents/{r.json()['document']['id']}", headers=auth).json()
    assert doc["status"] == "ready" and doc["chapter_count"] == 2


def test_upload_rejects_unsupported(client, auth):
    r = client.post("/api/documents/upload", headers=auth, files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 415


def test_sketches(client, auth):
    doc = client.post("/api/documents/text", headers=auth, json={"title": "Ink", "content": DESIGN_THINKING}).json()
    url = f"/api/documents/{doc['id']}/sketches"
    pts = [[0.1, 0.1, 0.5], [0.2, 0.15, 0.7]]
    s = client.post(url, headers=auth, json={"chapter_index": 1, "tool": "marker", "color": "#ffd60a", "width": 12, "points": pts})
    assert s.status_code == 201 and s.json()["points"] == pts
    client.post(url, headers=auth, json={"chapter_index": 2, "points": pts})
    assert len(client.get(url, headers=auth).json()) == 2
    assert len(client.get(url, headers=auth, params={"chapter": 1}).json()) == 1
    assert client.post(url, headers=auth, json={"chapter_index": 1, "color": "red", "points": pts}).status_code == 422
    assert client.delete(f"/api/documents/sketches/{s.json()['id']}", headers=auth).status_code == 204
    assert client.delete(url, headers=auth, params={"chapter": 2}).status_code == 204
    assert client.get(url, headers=auth).json() == []


def test_highlights_bookmarks_notes(client, auth):
    doc = client.post("/api/documents/text", headers=auth, json={"title": "Design", "content": DESIGN_THINKING}).json()
    h = client.post(f"/api/documents/{doc['id']}/highlights", headers=auth, json={"chapter_index": 1, "text": "affordance", "category": "Definition"})
    assert h.status_code == 201
    assert client.patch(f"/api/documents/highlights/{h.json()['id']}", headers=auth, json={"category": "Exam"}).json()["category"] == "Exam"
    b = client.post(f"/api/documents/{doc['id']}/bookmarks", headers=auth, json={"chapter_index": 2, "scroll_ratio": 0.3, "label": "Mapping"})
    assert b.status_code == 201
    assert len(client.get(f"/api/documents/{doc['id']}/bookmarks", headers=auth).json()) == 1
    n = client.post("/api/notes", headers=auth, json={"title": "Q", "content": "Why?", "document_id": doc["id"], "category": "Question"}).json()
    assert n["document_title"] == "Design"
    assert client.patch(f"/api/notes/{n['id']}", headers=auth, json={"pinned": True}).json()["pinned"]
    nb = client.get("/api/notes/notebook", headers=auth).json()
    assert nb["highlights"] and nb["notes"]
    assert "Study Notebook" in client.get("/api/notes/export/markdown", headers=auth).text
    assert client.delete(f"/api/notes/{n['id']}", headers=auth).status_code == 204


def test_ownership_isolation(client, auth, demo):
    mine = client.post("/api/documents/text", headers=auth, json={"title": "Private", "content": "Secret text " * 30}).json()
    assert client.get(f"/api/documents/{mine['id']}", headers=demo).status_code == 404


# ─── AI ───────────────────────────────────────────────────────────────────────
def test_ai_features(client, auth):
    doc = client.post("/api/documents/text", headers=auth, json={"title": "Focused Reading", "content": FOCUSED_READING}).json()
    did = doc["id"]
    for depth in ("brief", "standard", "deep"):
        s = client.post("/api/ai/summarize", headers=auth, json={"document_id": did, "chapter_index": 1, "depth": depth}).json()
        assert s["main_idea"] and s["key_points"] and s["concepts"]
    whole = client.post("/api/ai/summarize", headers=auth, json={"document_id": did, "depth": "standard"}).json()
    assert whole["scope"]["chapters"] == 4

    for mode in ("simple", "student", "expert", "example", "analogy"):
        e = client.post("/api/ai/explain", headers=auth, json={"text": "Recognition feels like knowledge, but only retrieval proves it.", "mode": mode, "document_id": did}).json()
        assert e["explanation"]

    chat = client.post("/api/ai/chat", headers=auth, json={"document_id": did, "question": "What is a mental model?"}).json()
    assert "mental model" in chat["answer"].lower()
    assert chat["citations"] and chat["citations"][0]["chapter_index"] == 1

    mm = client.post("/api/ai/mindmap", headers=auth, json={"document_id": did}).json()
    assert len(mm["root"]["children"]) == 4

    dg = client.post("/api/ai/diagram", headers=auth, json={"document_id": did, "chapter_index": 2, "kind": "flowchart"}).json()
    assert len(dg["nodes"]) >= 4 and dg["mermaid"].startswith("flowchart")
    cm = client.post("/api/ai/diagram", headers=auth, json={"document_id": did, "chapter_index": 1, "kind": "concept", "layout": "LR"}).json()
    assert cm["edges"]

    qz = client.post("/api/ai/quiz", headers=auth, json={"document_id": did, "chapter_index": 1, "count": 5}).json()
    mcq = [q for q in qz["questions"] if q["type"] == "multiple_choice"]
    assert mcq and all(0 <= q["correct"] < len(q["options"]) for q in mcq)

    g = client.post("/api/ai/quiz/grade", headers=auth, json={"answer": "a picture of how ideas relate", "reference": "a simplified internal picture of how the ideas in a text relate"}).json()
    assert g["score"] > 0.3

    cmp = client.post("/api/ai/compare", headers=auth, json={"document_id": did, "concept_a": "recognition", "concept_b": "retrieval"}).json()
    assert cmp["connection"]
    recs = client.get("/api/ai/recommendations", headers=auth).json()
    assert isinstance(recs, list)


# ─── Audio ────────────────────────────────────────────────────────────────────
def test_audiobook_generation(client, auth):
    doc = client.post("/api/documents/text", headers=auth, json={"title": "Listen", "content": FOCUSED_READING}).json()
    assert len(client.get("/api/audio/tones").json()) == 3
    r = client.post("/api/audio/audiobooks", headers=auth, json={"document_id": doc["id"], "tone": "Humor", "rate": 1.25})
    assert r.status_code == 202
    book = client.get(f"/api/audio/audiobooks/{r.json()['audiobook']['id']}", headers=auth).json()
    assert book["status"] == "ready" and len(book["script"]) == 4
    assert book["script"][0]["segments"][0]["kind"] == "intro"
    assert all(len(s["text"]) <= 240 for ch in book["script"] for s in ch["segments"])
    upd = client.patch(f"/api/audio/audiobooks/{book['id']}", headers=auth, json={"position_chapter": 2, "position_segment": 3}).json()
    assert upd["progress"] > 0
    assert "Humor" in client.get(f"/api/audio/audiobooks/{book['id']}/script.txt", headers=auth).text


def test_collab_studio(client, auth):
    doc = client.post("/api/documents/text", headers=auth, json={"title": "Collab", "content": DESIGN_THINKING}).json()
    p = client.post("/api/audio/collab", headers=auth, json={"title": "Team book", "document_id": doc["id"]}).json()
    assert len(p["assignments"]) == 4
    collabs = p["collaborators"] + [{"name": "Asha", "email": "asha@example.com", "role": "Narrator"}]
    p = client.patch(f"/api/audio/collab/{p['id']}", headers=auth, json={"collaborators": collabs}).json()
    asha = p["collaborators"][1]["id"]
    assigns = [{**a, "collaborator_id": asha, "status": "assigned"} if i == 0 else a for i, a in enumerate(p["assignments"])]
    p = client.patch(f"/api/audio/collab/{p['id']}", headers=auth, json={"assignments": assigns}).json()
    aid = p["assignments"][0]["id"]
    r = client.post(f"/api/audio/collab/{p['id']}/recordings", headers=auth,
                    data={"assignment_id": aid, "duration_seconds": "3.5"}, files={"file": ("take.webm", b"\x1aE\xdf\xa3fake", "audio/webm")})
    assert r.status_code == 201
    rec_id = r.json()["recording"]["id"]
    assert r.json()["project"]["assignments"][0]["status"] == "recorded"
    assert client.get(f"/api/audio/collab/{p['id']}/recordings/{rec_id}", headers=auth).content.startswith(b"\x1aE")


# ─── Research ─────────────────────────────────────────────────────────────────
def test_research_workflow(client, auth):
    src = client.post("/api/documents/text", headers=auth, json={"title": "Sound paper", "content": FOCUSED_READING, "doc_type": "PAPER"}).json()
    p = client.post("/api/research", headers=auth, json={"title": "Sound & focus", "topic": "ambient sound and reading focus", "source_document_ids": [src["id"]]}).json()
    disc = client.post(f"/api/research/{p['id']}/discover", headers=auth, json={}).json()
    assert disc["results"][0]["provider"] == "Your library"  # offline mode in tests
    refs = [
        {"title": "Background noise and reading comprehension", "authors": ["Jane Smith", "Raj Patel"], "year": 2021, "venue": "Journal of Reading",
         "abstract": "Steady background noise had little effect on reading comprehension, while intelligible speech reduced comprehension significantly.", "doi": "10.1000/x1"},
        {"title": "Music, lyrics and attention", "authors": ["Li Chen"], "year": 2019,
         "abstract": "Music with lyrics impaired sustained attention during demanding reading tasks compared with instrumental music."},
    ]
    added = client.post(f"/api/research/{p['id']}/references", headers=auth, json=refs).json()
    assert len(added["added"]) == 2
    assert client.post(f"/api/research/{p['id']}/references", headers=auth, json=refs).json()["added"] == []
    assert client.post(f"/api/research/{p['id']}/outline", headers=auth).json()["outline"]

    for key in ("abstract", "introduction", "literature_review", "methodology", "results", "conclusion"):
        sec = client.post(f"/api/research/{p['id']}/sections/{key}", headers=auth, json={"action": "generate"}).json()
        assert sec["section"]["content"].strip(), key
    for action in ("expand", "shorten", "academic", "cite"):
        sec = client.post(f"/api/research/{p['id']}/sections/introduction", headers=auth, json={"action": action}).json()
        assert sec["section"]["versions"]
    edited = client.put(f"/api/research/{p['id']}/sections/conclusion", headers=auth, json={"content": "My own conclusion [1]."}).json()
    assert edited["section"]["content"] == "My own conclusion [1]."
    restored = client.post(f"/api/research/{p['id']}/sections/conclusion/restore/0", headers=auth).json()
    assert restored["section"]["content"] != "My own conclusion [1]."

    for style in ("APA", "IEEE", "MLA"):
        cit = client.get(f"/api/research/{p['id']}/citations", headers=auth, params={"style": style}).json()
        assert len(cit["entries"]) == 2
    assert client.get(f"/api/research/{p['id']}/review", headers=auth).json()["disclaimer"]
    for fmt, sig in (("docx", b"PK"), ("pdf", b"%PDF"), ("md", b"# Sound")):
        r = client.post(f"/api/research/{p['id']}/export", headers=auth, json={"format": fmt})
        assert r.status_code == 200 and r.content.startswith(sig), fmt


# ─── Activity ─────────────────────────────────────────────────────────────────
def test_stats_and_search(client, demo):
    assert client.post("/api/sessions", headers=demo, json={"kind": "reading", "seconds": 600, "mode": "Zen"}).status_code == 201
    st = client.get("/api/stats", headers=demo).json()
    assert len(st["week"]) == 7 and st["streak"] >= 1 and st["today_minutes"] >= 10
    res = client.get("/api/search", headers=demo, params={"q": "affordance"}).json()
    assert res["passages"] and res["highlights"]
    assert client.get("/api/jobs", headers=demo).status_code == 200


# ─── Large uploads ────────────────────────────────────────────────────────────
def test_upload_limit_is_5gb_and_enforced_while_streaming(client, auth):
    from app.config import get_settings
    from app.services import storage

    settings = get_settings()
    assert settings.max_upload_mb == 5 * 1024
    original = settings.max_upload_mb
    settings.max_upload_mb = 1  # 1 MB for the test; the check runs chunk-by-chunk
    try:
        big = (b"word " * 300_000)  # ~1.5 MB
        r = client.post("/api/documents/upload", headers=auth, files={"file": ("big.txt", big, "text/plain")})
        assert r.status_code == 413, r.text
        leftovers = list(settings.storage_dir.rglob("*.part"))
        assert leftovers == [], leftovers  # partial file cleaned up
    finally:
        settings.max_upload_mb = original
    assert storage.CHUNK_BYTES >= 1024 * 1024


def test_streamed_upload_processes_from_disk(client, auth):
    body = ("# Part One\n\n" + "Streaming keeps memory flat for very large files. " * 4000).encode()
    r = client.post("/api/documents/upload", headers=auth, files={"file": ("stream.md", body, "text/markdown")})
    assert r.status_code == 202, r.text
    doc = client.get(f"/api/documents/{r.json()['document']['id']}", headers=auth).json()
    assert doc["status"] == "ready" and doc["file_size"] == len(body) and doc["word_count"] > 20000


# ─── Google sign-in ───────────────────────────────────────────────────────────
def test_google_signin(client, monkeypatch):
    from app.config import get_settings
    from app.services import google_auth

    assert client.get("/api/auth/config").json()["google_enabled"] is False
    assert client.post("/api/auth/google", json={"credential": "x" * 40}).status_code == 401  # not configured

    monkeypatch.setattr(get_settings(), "google_client_id", "test-client.apps.googleusercontent.com")
    claims = {"sub": "g-123", "email": "Reader.G@example.com", "email_verified": True, "name": "G Reader",
              "picture": "https://example.com/a.png", "iss": "accounts.google.com"}
    monkeypatch.setattr(google_auth, "verify_credential", lambda cred: claims)
    cfg = client.get("/api/auth/config").json()
    assert cfg["google_enabled"] and cfg["google_client_id"].endswith("googleusercontent.com")

    r = client.post("/api/auth/google", json={"credential": "y" * 40})
    assert r.status_code == 200, r.text
    u = r.json()["user"]
    assert u["email"] == "reader.g@example.com" and u["google_linked"] and not u["has_password"]
    h = {"Authorization": f"Bearer {r.json()['access_token']}"}
    # same Google account → same user; Google-only users can set a first password
    assert client.post("/api/auth/google", json={"credential": "y" * 40}).json()["user"]["id"] == u["id"]
    assert client.post("/api/auth/me/password", headers=h, json={"current_password": "", "new_password": "brandnew123"}).status_code == 204
    assert client.post("/api/auth/login", json={"email": "reader.g@example.com", "password": "brandnew123"}).status_code == 200

    # an existing password account with the same verified email is linked, not duplicated
    client.post("/api/auth/register", json={"name": "P", "email": "link@example.com", "password": "password123"})
    claims.update(sub="g-999", email="link@example.com")
    linked = client.post("/api/auth/google", json={"credential": "z" * 40}).json()["user"]
    assert linked["google_linked"] and linked["has_password"]

    def bad(cred):
        raise google_auth.GoogleAuthError("Google sign-in could not be verified. Please try again.")
    monkeypatch.setattr(google_auth, "verify_credential", bad)
    assert client.post("/api/auth/google", json={"credential": "q" * 40}).status_code == 401


def test_google_verification_rejects_forged_token(monkeypatch):
    from app.config import get_settings
    from app.services import google_auth

    monkeypatch.setattr(get_settings(), "google_client_id", "test-client.apps.googleusercontent.com")
    forged = "eyJhbGciOiJSUzI1NiIsImtpZCI6ImZha2UifQ.eyJzdWIiOiIxIn0.c2lnbmF0dXJl"
    try:
        google_auth.verify_credential(forged)
    except google_auth.GoogleAuthError:
        pass
    else:
        raise AssertionError("forged token was accepted")


def _pptx_bytes() -> bytes:
    """A minimal two-slide deck: slide 1 has a title, body and speaker notes; slide 2 only a title and a table."""
    import zipfile

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"

    def shape(text, ph=None):
        ph_xml = f'<p:nvPr><p:ph type="{ph}"/></p:nvPr>' if ph else "<p:nvPr/>"
        return f'<p:sp><p:nvSpPr><p:cNvPr id="1" name="s"/><p:cNvSpPr/>{ph_xml}</p:nvSpPr><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'

    def slide(*shapes, extra=""):
        return f'<p:sld xmlns:a="{a}" xmlns:p="{p}"><p:cSld><p:spTree>{"".join(shapes)}{extra}</p:spTree></p:cSld></p:sld>'

    table = '<a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>Cell A</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Cell B</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl>'
    notes = f'<p:notes xmlns:a="{a}" xmlns:p="{p}"><p:cSld><p:spTree>{shape("Remember to pause here.")}</p:spTree></p:cSld></p:notes>'
    rels = '<Relationships><Relationship Id="rId1" Target="../notesSlides/notesSlide1.xml"/></Relationships>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("ppt/slides/slide2.xml", slide(shape("Second slide", "title"), extra=table))
        zf.writestr("ppt/slides/slide1.xml", slide(shape("Opening slide", "title"), shape("Attention is a muscle " * 40)))
        zf.writestr("ppt/slides/_rels/slide1.xml.rels", rels)
        zf.writestr("ppt/notesSlides/notesSlide1.xml", notes)
    return buf.getvalue()


def test_extract_pptx_in_slide_order_with_notes_and_tables():
    from app.services import extraction

    doc = extraction.extract("deck.pptx", _pptx_bytes())
    text = "\n".join(c.title + "\n" + c.content for c in doc.chapters)
    assert text.index("Opening slide") < text.index("Second slide")  # numeric, not archive, order
    assert "Attention is a muscle" in text
    assert "Speaker notes: Remember to pause here." in text
    assert "Cell A | Cell B" in text


def test_upload_pptx(client, auth):
    r = client.post("/api/documents/upload", headers=auth, files={"file": ("deck.pptx", _pptx_bytes(), "application/octet-stream")})
    assert r.status_code == 202, r.text


def test_tts_aligns_words_and_caches(client, auth, monkeypatch):
    from app.services import tts

    calls = []

    async def fake_edge(text, voice, delivery):
        calls.append(voice)
        words = [(0, 300, "Hello"), (300, 250, "world,"), (550, 200, "again")]
        return b"ID3fake-mp3", tts._align(text, words), "ms"

    monkeypatch.setattr(tts, "_edge", fake_edge)
    assert client.get("/api/tts/status", headers=auth).json()["available"] is True
    assert any(v["id"].startswith("edge:en-IN") for v in client.get("/api/tts/voices", headers=auth).json())

    body = {"text": "Hello  world, again", "voice": "edge:en-IN-NeerjaNeural"}
    r = client.post("/api/tts/synthesize", headers=auth, json=body)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["mime"] == "audio/mpeg" and out["timing"] == "ms" and out["voice"] == "edge:en-IN-NeerjaNeural"
    assert out["words"] == [[0, 5, 0, 300], [6, 6, 300, 250], [13, 5, 550, 200]]  # offsets index the whitespace-normalised text

    assert client.post("/api/tts/synthesize", headers=auth, json=body).json()["words"] == out["words"]
    assert calls == ["en-IN-NeerjaNeural"]  # second request came from the disk cache


def test_tts_rejects_unknown_provider_and_requires_auth(client, auth):
    assert client.post("/api/tts/synthesize", json={"text": "hi"}).status_code == 401
    r = client.post("/api/tts/synthesize", headers=auth, json={"text": "hi", "voice": "nope:voice"})
    assert r.status_code == 502 and "isn’t available" in r.json()["detail"]


def test_hosted_postgres_urls_are_normalised():
    from app.database import _normalise_url

    assert _normalise_url("postgres://u:p@h:5432/db") == "postgresql+psycopg://u:p@h:5432/db"
    assert _normalise_url("postgresql://u:p@h/db") == "postgresql+psycopg://u:p@h/db"
    assert _normalise_url("sqlite:////data/x.db") == "sqlite:////data/x.db"
