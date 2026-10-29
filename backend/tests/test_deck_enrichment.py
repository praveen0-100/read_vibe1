"""Slides carry callouts (notations / exclamations), a main point, highlighted phrases and example slides with a generated scene."""
import io

from PIL import Image

from app.services import art, deck, llm

SECTIONS = [
    (0, "Password Security", (
        "A password is a secret string that proves who you are. Security means protecting accounts from attackers who guess or steal passwords. "
        "Surprisingly, most breaches start with a weak password! You should always use a long passphrase and a different one for every site. "
        "For example, a team that reuses one password across email and banking exposes every account when one site leaks it. "
        "Attackers try millions of common passwords every second, so length matters more than clever symbols.")),
    (1, "Responding to Incidents", (
        "An incident is any event that threatens the confidentiality of data. The first step is to contain the damage by disconnecting affected devices. "
        "Imagine a laptop stolen from a car: the team locks the account, wipes the device remotely and resets every password it stored. "
        "Afterwards write down what you learned so the next response is faster.")),
    (2, "Habits", "Remember that training is most effective when it is short and repeated. For instance, a monthly phishing drill changes behaviour more than a yearly lecture."),
]


def _local(monkeypatch):
    monkeypatch.setattr(llm, "available", lambda: False)  # deterministic: the local engine only
    return deck.build("Security at Work", SECTIONS, deck.DeckOptions(slides=14))["slides"]


def test_local_deck_has_callouts_main_points_and_examples(monkeypatch):
    slides = _local(monkeypatch)
    concepts = [s for s in slides if s["layout"] == "concept"]
    assert concepts and all(s["main_point"] for s in concepts)
    kinds = {c["kind"] for s in slides for c in s.get("callouts", [])}
    assert "wow" in kinds and kinds <= set(deck.CALLOUT_KINDS)
    wow = next(c for s in slides for c in s.get("callouts", []) if c["kind"] == "wow")
    assert wow["text"].endswith("!")                                                    # an exclamation reads like one
    for s in concepts:                                                                  # a main point is never repeated by a note
        assert all(c["text"].lower().rstrip("!.") != s["main_point"].lower().rstrip("!.") for c in s["callouts"])
        assert len(s["callouts"]) <= 2
        text = " ".join(s["bullets"]).lower()
        assert all(h.lower() in text for h in s["highlights"])                          # only phrases really on the slide

    examples = [s for s in slides if s["layout"] == "example"]
    assert examples and len(examples) <= 4
    for ex in examples:
        assert ex["data"]["text"] and ex["image"]["style"] == "scene" and ex["alt"]


def test_normalizer_cleans_callouts_and_highlights():
    raw = {"layout": "concept", "title": "Passwords", "bullets": ["Use a passphrase", "Never reuse one"],
           "main_point": "Length beats cleverness", "highlights": ["passphrase", "invisible phrase"],
           "callouts": [{"kind": "wow", "text": "Millions of guesses a second!"}, {"kind": "bogus", "text": "x"}, {"kind": "tip", "text": "a"}, {"kind": "note", "text": "b"}]}
    s = deck._normalize(raw, 0, deck.DeckOptions())
    assert s["highlights"] == ["passphrase"]                                            # a phrase that is not on the slide is dropped
    assert [c["kind"] for c in s["callouts"]] == ["wow", "tip"]                         # unknown kinds dropped, at most two kept
    ex = deck._normalize({"layout": "example", "title": "In practice", "data": {"text": "A team reuses a password.", "takeaway": "Reuse is risky"}}, 1, deck.DeckOptions())
    assert ex["layout"] == "example" and ex["image"]["style"] == "scene"
    broken = deck._normalize({"layout": "example", "title": "No text", "bullets": ["fallback"], "data": {}}, 2, deck.DeckOptions())
    assert broken["layout"] == "bullets"                                                # an example with nothing to show degrades to text


def test_scene_art_picks_a_pictogram_from_the_words_and_is_deterministic():
    assert art.pick_scene("A team of people in the community") == "people"
    assert art.pick_scene("passwords protect private data") == "shield"
    a = art.render("sales growth chart", "aurora", "scene", 480, 270)
    assert a == art.render("sales growth chart", "aurora", "scene", 480, 270)
    assert Image.open(io.BytesIO(a)).size == (480, 270)
    assert a != art.render("sales growth chart", "aurora", "orbit", 480, 270)           # the scene is not the abstract motif


def test_pptx_export_carries_the_new_elements(monkeypatch):
    from pptx import Presentation

    from app.services import pptx_export

    slides = _local(monkeypatch)
    data = pptx_export.build_pptx("Security at Work", "aurora", slides)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(slides)
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "KEY POINT" in text and "EXAMPLE" in text and "WOW!" in text
    pictures = sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13)
    assert pictures >= len([s for s in slides if s["layout"] == "example"])              # every example has its illustration


def test_highlights_match_whole_words_and_key_point_is_not_repeated(monkeypatch):
    from app.services import pptx_export

    assert pptx_export._runs("Designers care about design", ["design"]) == [("Designers care about ", False), ("design", True)]
    assert deck._highlights(["design"], "Designers only") == []
    for s in _local(monkeypatch):
        if s["layout"] == "concept" and s["main_point"]:
            assert s["main_point"].lower().rstrip(".!") not in [b.lower().rstrip(".!") for b in s["bullets"]] or len(s["bullets"]) <= 2
