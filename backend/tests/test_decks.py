import io

import pytest
from pptx import Presentation

from app.seed import DESIGN_THINKING
from app.services import art, deck as deck_svc

SAMPLE = """# Laying Plans
The art of war is of vital importance to the State. It is a matter of life and death, a road either to safety or to ruin. It is governed by five constant factors: the Moral Law, Heaven, Earth, the Commander, and Method and discipline. The Moral Law causes the people to be in complete accord with their ruler. Heaven signifies night and day, cold and heat, times and seasons.

# Waging War
First, count the cost. Then, gather the supplies. Next, move swiftly. Finally, secure the victory. There is no instance of a country having benefited from prolonged warfare. In 1999 and 2010 scholars noted that 70% of campaigns failed when delayed. The expenditure at home and at the front will reach a thousand ounces of silver per day.

# Attack By Stratagem
Supreme excellence consists in breaking the enemy's resistance without fighting. The highest form of generalship is to balk the enemy's plans; the next best is to prevent the junction of the enemy's forces. Thus the skillful leader subdues the enemy's troops without any fighting and takes cities without laying siege to them.
"""


def _wait_ready(client, auth, deck_id):
    d = client.get(f"/api/decks/{deck_id}", headers=auth).json()
    assert d["status"] == "ready", d
    return d


def test_deck_from_uploaded_text(client, auth):
    r = client.post("/api/decks/upload", headers=auth, files={"file": ("art_of_war.md", SAMPLE.encode(), "text/markdown")}, data={"slides": "10"})
    assert r.status_code == 202, r.text
    d = _wait_ready(client, auth, r.json()["deck"]["id"])
    layouts = [s["layout"] for s in d["slides"]]
    assert layouts[0] == "title" and layouts[-1] == "summary"
    assert {"concept", "mindmap"} <= set(layouts)
    assert d["meta"]["engine"] == "local"
    for s in d["slides"]:
        assert s["alt"] and s["notes"] and s["anim"] and s["transition"]
    assert d["title"] == "Art Of War"
    # switches are honoured when sent explicitly
    r2 = client.post("/api/decks/upload", headers=auth, files={"file": ("a.md", SAMPLE.encode(), "text/markdown")}, data={"visuals": "false", "notes": "false", "images": "false"})
    plain = _wait_ready(client, auth, r2.json()["deck"]["id"])
    assert not {"mindmap", "chart", "process"} & {s["layout"] for s in plain["slides"]}
    assert all(not s["notes"] and "image" not in s for s in plain["slides"])


def test_deck_accepts_xml_and_rejects_unknown_types(client, auth):
    xml = b"<book><title>Plants</title><section><title>Roots</title><p>Roots anchor a plant and absorb water and minerals from the soil. " * 5 + b"</p></section></book>"
    r = client.post("/api/decks/upload", headers=auth, files={"file": ("plants.xml", xml, "application/xml")})
    assert r.status_code == 202, r.text
    assert _wait_ready(client, auth, r.json()["deck"]["id"])["slide_count"] >= 3
    assert client.post("/api/decks/upload", headers=auth, files={"file": ("virus.exe", b"MZ", "application/octet-stream")}).status_code == 415
    assert client.post("/api/decks/upload", headers=auth, files={"file": ("empty.txt", b"", "text/plain")}).status_code == 400


def test_unreadable_upload_marks_deck_failed(client, auth):
    r = client.post("/api/decks/upload", headers=auth, files={"file": ("broken.pdf", b"not really a pdf", "application/pdf")})
    assert r.status_code == 202
    d = client.get(f"/api/decks/{r.json()['deck']['id']}", headers=auth).json()
    assert d["status"] == "failed" and d["meta"]["error"]
    assert client.get(f"/api/decks/{d['id']}/export.pptx", headers=auth).status_code == 409


def test_deck_from_library_document_edit_and_export(client, auth):
    up = client.post("/api/documents/upload", headers=auth, files={"file": ("dt.txt", DESIGN_THINKING.encode(), "text/plain")}).json()
    r = client.post("/api/decks", headers=auth, json={"document_id": up["document"]["id"] if "document" in up else up["id"], "slides": 14, "theme": "nebula"})
    assert r.status_code == 202, r.text
    d = _wait_ready(client, auth, r.json()["deck"]["id"])
    assert d["theme"] == "nebula" and d["meta"]["source"] == "library"

    # edits are re-validated: junk is dropped, a visual that lost its data degrades to text, the title persists
    slides = d["slides"]
    slides[1]["title"] = "Edited title"
    slides.append({"layout": "chart", "title": "Broken chart", "bullets": ["fallback text"], "data": {"labels": ["a"], "values": [1]}})
    slides.append({"layout": "bullets", "title": "", "bullets": ["no title → dropped"]})
    r = client.patch(f"/api/decks/{d['id']}", headers=auth, json={"slides": slides, "title": "My deck", "theme": "frost"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["title"] == "My deck" and out["theme"] == "frost" and out["slides"][1]["title"] == "Edited title"
    assert out["slides"][-1]["layout"] == "bullets" and out["slides"][-1]["title"] == "Broken chart"
    assert client.patch(f"/api/decks/{d['id']}", headers=auth, json={"slides": []}).status_code == 422

    pptx = client.get(f"/api/decks/{d['id']}/export.pptx", headers=auth)
    assert pptx.status_code == 200 and ".pptx" in pptx.headers["content-disposition"]
    prs = Presentation(io.BytesIO(pptx.content))
    assert len(prs.slides) == out["slide_count"]
    assert any(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in prs.slides)
    assert any(sh.has_chart for s in prs.slides for sh in s.shapes if hasattr(sh, "has_chart")) or all(s["layout"] != "chart" for s in out["slides"])

    assert any(x["id"] == d["id"] for x in client.get("/api/decks", headers=auth).json())
    assert client.delete(f"/api/decks/{d['id']}", headers=auth).status_code == 204
    assert client.get(f"/api/decks/{d['id']}", headers=auth).status_code == 404


def test_deck_ownership(client, auth, demo):
    r = client.post("/api/decks/upload", headers=auth, files={"file": ("mine.txt", SAMPLE.encode(), "text/plain")})
    deck_id = r.json()["deck"]["id"]
    assert client.get(f"/api/decks/{deck_id}", headers=demo).status_code == 404
    assert client.delete(f"/api/decks/{deck_id}", headers=demo).status_code == 404
    assert client.get("/api/decks").status_code == 401


def test_art_endpoint_is_deterministic_and_cached(client):
    a = client.get("/api/decks/art/image", params={"seed": "ocean currents", "theme": "solar", "w": 480, "h": 270})
    b = client.get("/api/decks/art/image", params={"seed": "ocean currents", "theme": "solar", "w": 480, "h": 270})
    c = client.get("/api/decks/art/image", params={"seed": "mountain", "theme": "solar", "w": 480, "h": 270})
    assert a.status_code == 200 and a.headers["content-type"] == "image/jpeg" and "immutable" in a.headers["cache-control"]
    assert a.content == b.content != c.content
    assert client.get("/api/decks/art/image", params={"seed": "x", "w": 99999}).status_code == 422


@pytest.mark.parametrize("theme", list(art.THEMES))
def test_every_theme_renders_every_style(theme):
    for style in art.STYLES:
        assert art.render("seed", theme, style, 320, 180)[:2] == b"\xff\xd8"


def test_normalizer_repairs_model_output():
    opts = deck_svc.DeckOptions(slides=8)
    ok = deck_svc._normalize({"layout": "chart", "title": "Sales", "data": {"kind": "bar", "labels": ["Q1", "Q2"], "values": ["10", 20]}}, 0, opts)
    assert ok["layout"] == "chart" and ok["data"]["values"] == [10.0, 20.0] and "Q1 10" in ok["alt"]
    assert deck_svc._normalize({"layout": "process", "title": "Two steps", "bullets": ["a"], "data": {"nodes": [{"label": "x"}, {"label": "y"}]}}, 0, opts)["layout"] == "bullets"
    assert deck_svc._normalize({"layout": "nonsense", "title": "T", "bullets": ["b"]}, 0, opts)["layout"] == "bullets"
    assert deck_svc._normalize({"layout": "concept", "title": "No bullets"}, 0, opts) is None
    plain = deck_svc._normalize({"layout": "mindmap", "title": "Map", "data": {"root": "r", "branches": [{"label": str(i), "children": ["c"]} for i in range(4)]}}, 0, deck_svc.DeckOptions(visuals=False))
    assert plain["layout"] == "bullets" and plain["bullets"] == ["0", "1", "2", "3"]
